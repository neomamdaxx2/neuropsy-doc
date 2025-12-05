from docx import Document
import yaml
import os

# Create directories if not exist
os.makedirs('templates', exist_ok=True)
os.makedirs('data', exist_ok=True)

# 1. Create Word Template
doc = Document()
doc.add_heading('Facture pour {{ client_name }}', 0)

p = doc.add_paragraph('Date : {{ date }}')
p = doc.add_paragraph('Adresse : {{ address }}')

doc.add_heading('Détails', level=1)

table = doc.add_table(rows=1, cols=3)
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Item'
hdr_cells[1].text = 'Quantité'
hdr_cells[2].text = 'Prix'

row_cells = table.add_row().cells
row_cells[0].text = '{{ item_1_name }}'
row_cells[1].text = '{{ item_1_qty }}'
row_cells[2].text = '{{ item_1_price }}'

row_cells = table.add_row().cells
row_cells[0].text = '{{ item_2_name }}'
row_cells[1].text = '{{ item_2_qty }}'
row_cells[2].text = '{{ item_2_price }}'

doc.add_paragraph('\nTotal: {{ total_price }} €')

template_path = os.path.join('templates', 'exemple.docx')
doc.save(template_path)
print(f"Created {template_path}")

# 2. Create YAML Data
data = {
    'client_name': 'Jean Dupont',
    'date': '05/12/2025',
    'address': '123 Rue de la Paix, Paris',
    'item_1_name': 'Consultation',
    'item_1_qty': 1,
    'item_1_price': 50,
    'item_2_name': 'Rapport écrit',
    'item_2_qty': 1,
    'item_2_price': 30,
    'total_price': 80
}

with open('data/exemple.yaml', 'w', encoding='utf-8') as f:
    yaml.dump(data, f, allow_unicode=True)

print("Created data/exemple.yaml")
