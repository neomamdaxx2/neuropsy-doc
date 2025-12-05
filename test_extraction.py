from generator import get_identifiers_from_template
from docx import Document
import os

# Create a test template
test_template_path = 'templates/extraction_test.docx'
doc = Document()
doc.add_paragraph('Voici une variable : {{ var1 }}')
doc.add_paragraph('Une autre : {{ var2 }}')
doc.add_paragraph('Et une répétée : {{ var1 }}')
doc.save(test_template_path)

print(f"Template created at {test_template_path}")

# Test extraction
try:
    keys = get_identifiers_from_template(test_template_path)
    print(f"Extracted keys: {keys}")
    
    expected_keys = {'var1', 'var2'}
    if keys == expected_keys:
        print("SUCCESS: Extraction verified.")
    else:
        print(f"FAILURE: Expected {expected_keys}, got {keys}")
        
except Exception as e:
    print(f"ERROR: {e}")
